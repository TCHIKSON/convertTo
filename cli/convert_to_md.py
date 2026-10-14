#!/usr/bin/env python3
"""
convert_to_md.py — Convertit des fichiers PDF, Word (.docx) et Excel (.xlsx/.xlsm) en Markdown.

Installation des dépendances :
    pip install pdfplumber python-docx openpyxl

Utilisation :
    python convert_to_md.py fichier.pdf
    python convert_to_md.py rapport.docx tableau.xlsx
    python convert_to_md.py dossier/              # convertit tous les fichiers pris en charge
    python convert_to_md.py dossier/ -r -o sortie/ # récursif, résultats dans sortie/
"""

import argparse
import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm"}


# ---------------------------------------------------------------- utilitaires
def _cell(value) -> str:
    """Nettoie une valeur pour l'insérer dans une cellule de tableau Markdown."""
    if value is None:
        return ""
    text = str(value).strip()
    return text.replace("|", "\\|").replace("\r", " ").replace("\n", "<br>")


def table_to_md(rows) -> str:
    """Transforme une liste de lignes (listes de valeurs) en tableau Markdown."""
    rows = [[_cell(c) for c in row] for row in rows if row is not None]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


# ---------------------------------------------------------------------- PDF
def pdf_to_md(path: Path) -> str:
    import pdfplumber

    parts = [f"# {path.stem}\n"]
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            parts.append(f"\n## Page {i}\n")
            tables = page.find_tables()
            # Texte hors des zones de tableau, pour ne pas dupliquer le contenu
            text_page = page
            for t in tables:
                text_page = text_page.outside_bbox(t.bbox)
            text = (text_page.extract_text() or "").strip()
            if text:
                parts.append(text + "\n")
            for t in tables:
                md = table_to_md(t.extract())
                if md:
                    parts.append("\n" + md + "\n")
            if not text and not tables:
                parts.append("_(page sans texte extractible — probablement une image scannée)_\n")
    return "\n".join(parts)


# --------------------------------------------------------------------- Word
def _runs_to_md(paragraph) -> str:
    out = []
    for run in paragraph.runs:
        t = run.text
        if not t:
            continue
        stripped = t.strip()
        if stripped and (run.bold or run.italic):
            lead = t[: len(t) - len(t.lstrip())]
            trail = t[len(t.rstrip()):]
            if run.bold and run.italic:
                stripped = f"***{stripped}***"
            elif run.bold:
                stripped = f"**{stripped}**"
            else:
                stripped = f"*{stripped}*"
            t = lead + stripped + trail
        out.append(t)
    return "".join(out).strip()


def docx_to_md(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    parts = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = _runs_to_md(p)
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style == "title":
                parts.append(f"# {p.text.strip()}")
            elif style.startswith("heading") or style.startswith("titre"):
                digits = "".join(c for c in style if c.isdigit())
                level = min(int(digits) if digits else 1, 6)
                parts.append(f"{'#' * level} {p.text.strip()}")
            elif "list number" in style or "liste num" in style:
                parts.append(f"1. {text}")
            elif "list" in style or "liste" in style or p._p.pPr is not None and p._p.pPr.numPr is not None:
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif tag == "tbl":
            t = Table(child, doc)
            rows = []
            for row in t.rows:
                cells, seen = [], set()
                for c in row.cells:  # évite les doublons des cellules fusionnées
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text)
                rows.append(cells)
            md = table_to_md(rows)
            if md:
                parts.append(md)
    # Regroupe les éléments de liste consécutifs sans ligne vide entre eux
    out = []
    for i, block in enumerate(parts):
        is_item = block.startswith(("- ", "1. "))
        prev_item = i > 0 and parts[i - 1].startswith(("- ", "1. "))
        out.append(("\n" if is_item and prev_item else "\n\n") + block if i else block)
    return "".join(out) + "\n"


# -------------------------------------------------------------------- Excel
def xlsx_to_md(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    parts = [f"# {path.stem}\n"]
    for ws in wb.worksheets:
        parts.append(f"\n## {ws.title}\n")
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # Supprime les colonnes entièrement vides
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [None] * (width - len(r)) for r in rows]
            keep = [j for j in range(width) if any(r[j] not in (None, "") for r in rows)]
            rows = [[r[j] for j in keep] for r in rows]
        md = table_to_md(rows)
        parts.append((md or "_(feuille vide)_") + "\n")
    wb.close()
    return "\n".join(parts)


# --------------------------------------------------------------------- main
CONVERTERS = {".pdf": pdf_to_md, ".docx": docx_to_md, ".xlsx": xlsx_to_md, ".xlsm": xlsx_to_md}


def collect(inputs, recursive):
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            pattern = "**/*" if recursive else "*"
            for f in sorted(p.glob(pattern)):
                if f.is_file() and f.suffix.lower() in SUPPORTED and not f.name.startswith("~$"):
                    yield f
        elif p.is_file():
            yield p
        else:
            print(f"⚠️  Introuvable : {p}", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description="Convertit PDF, Word (.docx) et Excel (.xlsx) en Markdown.")
    ap.add_argument("inputs", nargs="+", help="fichiers ou dossiers à convertir")
    ap.add_argument("-o", "--output", help="dossier de sortie (par défaut : à côté du fichier source)")
    ap.add_argument("-r", "--recursive", action="store_true", help="parcourir les sous-dossiers")
    args = ap.parse_args()

    out_dir = Path(args.output) if args.output else None
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)

    ok = fail = 0
    for f in collect(args.inputs, args.recursive):
        ext = f.suffix.lower()
        if ext not in CONVERTERS:
            hint = " (enregistrez-le d'abord en .docx / .xlsx)" if ext in {".doc", ".xls"} else ""
            print(f"⏭️  Format non pris en charge : {f.name}{hint}", file=sys.stderr)
            continue
        target = (out_dir or f.parent) / (f.stem + ".md")
        try:
            target.write_text(CONVERTERS[ext](f), encoding="utf-8")
            print(f"✅ {f.name} → {target}")
            ok += 1
        except Exception as e:  # noqa: BLE001
            print(f"❌ {f.name} : {e}", file=sys.stderr)
            fail += 1
    print(f"\nTerminé : {ok} converti(s), {fail} erreur(s).")
    sys.exit(1 if fail else 0)


if __name__ == "__main__":
    main()